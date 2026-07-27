"""
Paso 1 — Fuente de datos (Excel / Google Sheets)
Proyecto de final de curso — Videojuego (seed42)

El alumno parte de dos archivos CSV con los datos en bruto (partidas y
jugadores). Este módulo documenta cada columna (nombre, tipo esperado,
valores posibles) y genera un Excel con esa documentación + los datos en
bruto, para "leer la realidad de los datos antes de tocarlos con código".
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from lib.resumen import registrar


# Diccionario de datos: para cada columna, (tipo esperado, valores posibles/notas).
# Se rellena a mano tras inspeccionar los CSV con pandas (df.dtypes, df["col"].unique()...).
DICCIONARIO_PARTIDAS = [
    ("id_partida", "Entero", "Identificador único de la partida. Puede tener duplicados a limpiar."),
    ("fecha", "Fecha", "Mezcla 3 formatos: AAAA-MM-DD, DD-MM-AA, DD/MM/AAAA."),
    ("hora", "Texto (HH:MM)", "Hora de inicio de la partida."),
    ("id_jugador", "Texto (código)", "Formato 'J'+4 dígitos (J0001-J0200). Relaciona con la tabla jugadores."),
    ("mapa", "Categórico", "Mapa jugado. 6 valores, sin inconsistencias detectadas."),
    ("modo", "Categórico", "Modo de juego: Casual, Ranked o Torneo."),
    ("personajes_usados", "Texto (multivalor)", "Personajes separados por ';'. Mayúsculas inconsistentes. Contiene nulos."),
    ("duracion_min", "Entero", "Duración de la partida en minutos."),
    ("kills", "Decimal", "Bajas conseguidas. Contiene nulos."),
    ("muertes", "Entero", "Veces que murió el jugador."),
    ("asistencias", "Entero", "Asistencias en la partida."),
    ("resultado", "Categórico", "Victoria o Derrota. Mayúsculas/espacios inconsistentes. Contiene nulos."),
    ("puntos", "Decimal (texto a convertir)", "Algunos valores usan coma decimal y sufijo ' pts'. Contiene atípicos."),
]

DICCIONARIO_JUGADORES = [
    ("id_jugador", "Texto (código)", "Identificador único, formato 'J'+4 dígitos."),
    ("nick", "Texto", "Nombre de usuario del jugador."),
    ("region", "Categórico", "Región de juego: EU-Oeste, EU-Este, LATAM, ASIA. Contiene nulos."),
    ("nivel", "Entero", "Nivel de experiencia del jugador."),
    ("rango", "Categórico", "Rango competitivo: Bronce, Plata, Oro, Platino, Diamante."),
    ("fecha_registro", "Fecha", "Fecha de alta del jugador, formato AAAA-MM-DD."),
]


def cargar_datos(ruta_partidas: str, ruta_jugadores: str) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Lee los dos CSV en bruto, tal cual vienen, sin transformar nada."""
    df_partidas = pd.read_csv(ruta_partidas)
    df_jugadores = pd.read_csv(ruta_jugadores)
    return df_partidas, df_jugadores


def _escribir_bloque_diccionario(hoja, fila_inicio: int, titulo: str, filas: list[tuple]) -> int:
    """Escribe un bloque (título + tabla de 3 columnas) en la hoja de diccionario,
    empezando en fila_inicio. Devuelve la siguiente fila libre para el próximo bloque."""
    fuente_titulo = Font(name="Arial", bold=True, size=12, color="305496")
    fuente_cabecera = Font(name="Arial", bold=True, color="FFFFFF")
    relleno_cabecera = PatternFill(start_color="305496", fill_type="solid")

    hoja.cell(row=fila_inicio, column=1, value=titulo).font = fuente_titulo
    fila = fila_inicio + 1
    for i, texto in enumerate(["Columna", "Tipo esperado", "Valores posibles / notas"], start=1):
        celda = hoja.cell(row=fila, column=i, value=texto)
        celda.font = fuente_cabecera
        celda.fill = relleno_cabecera
    fila += 1
    for nombre, tipo, nota in filas:
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=tipo)
        hoja.cell(row=fila, column=3, value=nota).alignment = Alignment(wrap_text=True, vertical="top")
        fila += 1
    return fila + 1  # deja una fila en blanco antes del siguiente bloque


def generar_excel_bruto(df_partidas: pd.DataFrame, df_jugadores: pd.DataFrame, ruta_salida: str) -> None:
    """
    Genera un Excel con 3 hojas a partir de los DataFrames en bruto:
      - diccionario_datos: documentación de cada columna
      - partidas: datos en bruto de las partidas
      - jugadores: datos en bruto de los jugadores
    Este es el entregable del Paso 1, antes de limpiar nada.
    """
    with pd.ExcelWriter(ruta_salida, engine="openpyxl") as writer:
        df_partidas.to_excel(writer, sheet_name="partidas", index=False)
        df_jugadores.to_excel(writer, sheet_name="jugadores", index=False)

    libro = load_workbook(ruta_salida)
    hoja = libro.create_sheet("diccionario_datos", 0)
    siguiente_fila = _escribir_bloque_diccionario(
        hoja, 1, f"Hoja: partidas ({len(df_partidas)} filas)", DICCIONARIO_PARTIDAS
    )
    _escribir_bloque_diccionario(hoja, siguiente_fila, f"Hoja: jugadores ({len(df_jugadores)} filas)", DICCIONARIO_JUGADORES)
    hoja.column_dimensions["A"].width = 20
    hoja.column_dimensions["B"].width = 26
    hoja.column_dimensions["C"].width = 90

    # ajusta el ancho de columnas y congela la cabecera en las hojas de datos
    for nombre_hoja in ["partidas", "jugadores"]:
        hoja_datos = libro[nombre_hoja]
        for columna in hoja_datos.columns:
            ancho = max((len(str(c.value)) for c in columna if c.value is not None), default=10)
            hoja_datos.column_dimensions[get_column_letter(columna[0].column)].width = min(ancho + 2, 40)
        hoja_datos.freeze_panes = "A2"

    libro.save(ruta_salida)
    registrar("Paso 1", f"Excel generado: {ruta_salida}")
